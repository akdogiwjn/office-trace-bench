import sys
import subprocess
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import office
from office_trace_bench.contracts import ROOT,load_manifest,read_json,sha256
from office_trace_bench import records
from scripts.freeze_suite import lines,table_rows,chart_oracle
from scripts.native_xlsx_oracle import calculate
from openpyxl import load_workbook

class FrozenSuiteTests(unittest.TestCase):
    def test_exactly_seven_inputs_with_frozen_provenance(self):
        paths=list((ROOT/'datasets').glob('*/*/manifest.json'))
        self.assertEqual(len(paths),7)
        for path in paths:
            m=load_manifest(path);p=read_json(path.parent/'provenance.json')
            input_path=path.parent/'input'/m['workbook' if m['kind']=='xlsx' else 'form']
            self.assertEqual(sha256(input_path),p['sha256'])
            if p.get('source_evidence') and p.get('source_evidence_sha256'):
                self.assertEqual(sha256(ROOT/p['source_evidence']),p['source_evidence_sha256'])
            if m['kind']=='xlsx':
                meta=read_json(path.parent/'complexity.json');self.assertEqual(meta['input_sha256'],p['sha256'])
                self.assertFalse(p['synthetic_expansion'])
                oracle = (calculate(m['dataset_id'],input_path)['chart_values'] if p.get('input_class')=='official_native_spreadsheet'
                          else chart_oracle(m['dataset_id'],input_path))
                self.assertEqual(oracle,read_json(path.parent/'expected.json')['chart_values'])
                serialized=str(m['requirements'])
                for forbidden in ['required_references','selector_cell','formula_metrics','base_sheets','raw_sheets']:
                    self.assertNotIn(forbidden,serialized)
                if p.get('input_class')=='official_table_transcription':
                    retail=m['dataset_id']=='retail'
                    rows=table_rows(lines(ROOT/p['source_evidence'],4 if retail else 0),147 if retail else 8,192 if retail else 90,12 if retail else 13)
                    book=load_workbook(input_path,read_only=True,data_only=True)
                    try:
                        actual=[list(row) for row in book.worksheets[0].iter_rows(min_row=2,max_row=len(rows)+1,values_only=True)]
                        self.assertEqual(actual,rows);self.assertEqual(len(rows),p['observations'])
                    finally:book.close()
    def test_native_download_evidence_and_unchanged_hr_trace(self):
        evidence=read_json(ROOT/'reports/official-native-input-provenance-v3.json')
        expected_totals={'retail':('adjusted_sales',729538,35),'manufacturing':('total_shipments',601257,8)}
        for name,(metric,value,sheets) in expected_totals.items():
            root=ROOT/'datasets/xlsx'/name;m=load_manifest(root/'manifest.json');p=read_json(root/'provenance.json')
            self.assertEqual(p['input_class'],'official_native_spreadsheet')
            d=evidence['datasets'][name]
            self.assertTrue(d['byte_identity_verified']);self.assertEqual(d['official_redownload_sha256'],p['sha256'])
            self.assertEqual(d['local_sha256'],sha256(root/'input'/m['workbook']))
            self.assertEqual(calculate(name,root/'input'/m['workbook']),read_json(root/'expected.json'))
            self.assertEqual(read_json(root/'expected.json')['metric_values'][metric],value)
            self.assertEqual(read_json(root/'complexity.json')['sheet_count'],sheets)
        root=ROOT/'datasets/xlsx/hr';s=read_json(root/'provenance-supplement.json')
        self.assertEqual(s['manifest_sha256'],sha256(root/'manifest.json'))
        self.assertEqual(s['original_provenance_sha256'],sha256(root/'provenance.json'))
        self.assertEqual(s['input_sha256'],evidence['datasets']['hr']['local_sha256'])
        self.assertEqual(s['evidence_sha256'],sha256(ROOT/s['evidence_path']))
        self.assertFalse(s['independent_official_redownload'])
    def test_pdf_records_reproduce_and_hashes_are_frozen(self):
        paths=[p for p in (ROOT/'datasets').rglob('*') if p.is_file() and '__pycache__' not in p.parts]
        before={str(p):sha256(p) for p in paths}
        for args in [('scripts/freeze_pdf_metadata.py',),('scripts/freeze_suite.py','--initialize-v2'),
                     ('scripts/adopt_native_xlsx.py','--input-dir','missing-source','--official-redownload-dir','missing-source','--download-date','2026-10-09')]:
            process=subprocess.run([sys.executable,*args],cwd=ROOT,capture_output=True,text=True)
            self.assertNotEqual(process.returncode,0);self.assertIn('already frozen',process.stderr)
        self.assertEqual(before,{str(p):sha256(p) for p in paths})
        for dataset,fn in [('irs_w4',records.w4),('sba1919',records.sba)]:
            root=ROOT/'datasets/pdf'/dataset;m=load_manifest(root/'manifest.json');p=root/'input'/m['records'];contract=read_json(root/'records_contract.json')
            self.assertEqual(read_json(p),fn());self.assertEqual(sha256(p),contract['records_sha256']);self.assertEqual(contract['record_count'],m['requirements']['record_count'])
            self.assertEqual(sha256(ROOT/contract['generator']),contract['generator_sha256'])
        source=ROOT/'legacy/runner/cases/pdf/input/synthetic_applicants.json'
        self.assertEqual(sha256(source),sha256(ROOT/'datasets/pdf/opm/input/synthetic_applicants.json'))
        legacy=read_json(ROOT/'legacy/runner/cases/xlsx/expected.json')['cached_values']
        current=read_json(ROOT/'datasets/xlsx/tlc/expected.json')['metric_values']
        for name,address in [('total_trips','B5'),('fare_revenue','B6'),('average_fare','B7'),('removal_rate','B8'),('reconciliation_status','B9')]:
            self.assertEqual(current[name],legacy[address])
