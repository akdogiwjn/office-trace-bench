"""Real workbook mutations test preservation and flexible Agent-chosen layouts."""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import sys
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import office
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,LineChart,Reference
from openpyxl.styles import Font
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from office_trace_bench.contracts import write_json,sha256
from office_trace_bench.verify import verify,Checks
from office_trace_bench.spreadsheet import original_content
from office_trace_bench.complexity import inspect_workbook


def caches(path,values):
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    # Added Overview is appended after the original Invoices sheet.
    tree=ET.fromstring(parts['xl/worksheets/sheet2.xml'])
    for node in tree.findall('.//s:c',ns):
        if node.attrib['r'] in values:
            v=node.find('s:v',ns)
            if v is None:v=ET.SubElement(node,'{'+ns['s']+'}v')
            v.text=str(values[node.attrib['r']])
    parts['xl/worksheets/sheet2.xml']=ET.tostring(tree)
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for n,data in parts.items():z.writestr(n,data)


class SemanticSpreadsheetTests(unittest.TestCase):
    def fixture(self,root,cell='D9'):
        (root/'input').mkdir();(root/'output').mkdir()
        b=Workbook();s=b.active;s.title='Invoices';s.append(['Region','Amount']);s.append(['East',4]);s.append(['West',6]);s['D1']='=SUM(B2:B3)'
        s['A1'].font=Font(bold=True);s.merge_cells('E1:F1');s['E1']='Source'
        original=BarChart();original.add_data(Reference(s,min_col=2,min_row=1,max_row=3),titles_from_data=True);s.add_chart(original,'H2')
        seed=root/'input/seed.xlsx';b.save(seed)
        s=b.create_sheet('Agent Overview');s[cell]="=SUM(Invoices!B2:B3)";s[cell].font=Font(color='008000')
        s['C20']='Base';s['C20'].font=Font(color='0000FF')
        s['G5']='Base';s['G6']='Growth';s['H5']=1;s['H6']=2
        for name in ['H5','H6']:s[name].font=Font(color='0000FF')
        s['F22']=f'={cell}*INDEX(H5:H6,MATCH(C20,G5:G6,0))'
        s['L2']='East';s['L3']='West';s['M2']='=Invoices!B2';s['M3']='=Invoices!B3'
        chart=BarChart();chart.add_data(Reference(s,min_col=13,min_row=2,max_row=3));chart.set_categories(Reference(s,min_col=12,min_row=2,max_row=3));s.add_chart(chart,'P1')
        v=DataValidation(type='list',formula1='"Base,Growth"');s.add_data_validation(v);v.add(s['C20']);s.freeze_panes='B4'
        s.conditional_formatting.add(cell,CellIsRule(operator='greaterThan',formula=['0']));s[cell].comment=Comment('source','fixture')
        output=root/'output/result.xlsx';b.save(output);b.close();caches(output,{cell:10,'F22':10,'M2':4,'M3':6})
        loc=lambda a:dict(sheet='Agent Overview',cell=a)
        write_json(root/'output/workbook_features.json',dict(summary_sheet='Agent Overview',metrics=dict(total=loc(cell)),
            scenario=dict(selector=loc('C20'),assumptions=dict(Base=dict(scale=loc('H5')),Growth=dict(scale=loc('H6'))),projections=dict(projected=loc('F22'))),
            charts=dict(regions=dict(sheet='Agent Overview',index=0))))
        write_json(root/'expected.json',dict(metric_values=dict(total=10),scenario_values=dict(Base=dict(projected=10),Growth=dict(projected=20)),chart_values=dict(regions=[4,6])))
        write_json(root/'output/formula_recalc.json',dict(status='success',total_errors=0,total_formulas=5));write_json(root/'output/xlsx_enhancement_summary.json',dict(final_verifier_status='success'))
        manifest=dict(schema_version='office-dataset-v2',workload_contract='spreadsheet-semantic-v2',kind='xlsx',dataset_id='fixture',task='Invoices',
            workbook='seed.xlsx',output_workbook='result.xlsx',input_files={'input/seed.xlsx':sha256(seed)},expected_file='expected.json',expected_sha256=sha256(root/'expected.json'),trace_rules=[],
            requirements=dict(metrics=[dict(id='total',meaning='invoice total',unit='USD')],scenario=dict(default='Base',rows=[dict(name='Base',assumptions=dict(scale=1)),dict(name='Growth',assumptions=dict(scale=2))],projections=[dict(id='projected')]),
                              charts=[dict(id='regions')],minimum_conditional_formats=1,minimum_comments=1),
            required_outputs=['result.xlsx','workbook_features.json','formula_recalc.json','xlsx_enhancement_summary.json'])
        write_json(root/'manifest.json',manifest)
        return root/'manifest.json',output

    def test_different_metric_addresses_and_natural_sheet_counts_pass(self):
        for address in ['D9','Z18']:
            with TemporaryDirectory() as tmp:
                root=Path(tmp);manifest,_=self.fixture(root,address)
                self.assertEqual(verify(manifest,root)['failures'],[])
                meta=inspect_workbook(root/'input/seed.xlsx')
                self.assertEqual(meta['sheet_count'],1);self.assertEqual(meta['formula_count'],1);self.assertEqual(meta['chart_count'],1);self.assertEqual(meta['merged_range_count'],1)

    def test_source_mutations_and_missing_objects_fail(self):
        for mutation,check in [('data','source:data:Invoices'),('blank','source:data:Invoices'),('expansion','source:data:Invoices'),('formula','source:formulas:Invoices'),('sheet','source:worksheets'),('chart','source:charts:Invoices'),('chart_type','source:charts:Invoices'),('style','source:styles:Invoices'),('default_style','source:styles:Invoices'),('merge','source:merged_regions:Invoices')]:
            with TemporaryDirectory() as tmp:
                root=Path(tmp);manifest,output=self.fixture(root);b=load_workbook(output)
                if mutation=='data':b['Invoices']['B2']=99
                elif mutation=='blank':b['Invoices']['C2']=999
                elif mutation=='expansion':b['Invoices'].append(['fabricated row',999])
                elif mutation=='formula':b['Invoices']['D1']=10
                elif mutation=='sheet':del b['Invoices']
                elif mutation=='chart':b['Invoices']._charts=[]
                elif mutation=='chart_type':
                    chart=LineChart();chart.series=b['Invoices']._charts[0].series;b['Invoices']._charts=[chart]
                elif mutation=='style':b['Invoices']['A1'].font=Font(bold=False)
                elif mutation=='default_style':b['Invoices']['B2'].font=Font(italic=True)
                elif mutation=='merge':b['Invoices'].unmerge_cells('E1:F1')
                b.save(output);b.close();report=verify(manifest,root)
                self.assertIn(check,report['failures'],report)

    def test_hardcoded_answer_and_disconnected_scenario_fail(self):
        for mutation,wanted in [('metric','feature:formula:Agent Overview!D9'),('scenario','semantic:scenario_dependency:projected'),('chart','semantic:chart:regions')]:
            with TemporaryDirectory() as tmp:
                root=Path(tmp);manifest,output=self.fixture(root);b=load_workbook(output)
                if mutation=='metric':b['Agent Overview']['D9']=10
                elif mutation=='scenario':b['Agent Overview']['F22']='=D9*1'
                elif mutation=='chart':b['Agent Overview']['M2']='=Invoices!B3'
                b.save(output);b.close();caches(output,{'D9':10,'F22':10,'M2':6 if mutation=='chart' else 4,'M3':6})
                self.assertIn(wanted,verify(manifest,root)['failures'])
