#!/usr/bin/env python3
"""Trusted positive fixture for the verifier; never staged into an Agent run."""
import argparse
from pathlib import Path
import shutil
import subprocess
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import office
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from office_trace_bench.contracts import dataset_path, read_json, write_json
from office_trace_bench.runner import stage
from office_trace_bench.verify import verify


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--dataset',required=True);p.add_argument('--workspace',type=Path,required=True)
    args=p.parse_args();workspace=args.workspace.resolve();path=dataset_path('xlsx',args.dataset)
    manifest=stage(path,workspace);req=manifest['requirements'];output=workspace/'output'
    book=load_workbook(workspace/'input'/manifest['workbook'])
    sheet=book.create_sheet(req['summary_sheet'],0);sheet.freeze_panes='A5'
    for m in req['metrics']:
        cell=sheet[m['cell']];cell.value='='+m['required_references'][0]
        cell.font=Font(name='Arial',color=m.get('font_color','008000'))
        if m.get('number_format'):cell.number_format=m['number_format']
        sheet['A'+str(cell.row)]=m['name']
    scenario=req['scenario'];selector=scenario['selector_cell']
    sheet[selector]=scenario['default'];sheet[selector].font=Font(name='Arial',color='0000FF')
    dv=DataValidation(type='list',formula1='"'+','.join(x['name'] for x in scenario['rows'])+'"')
    sheet.add_data_validation(dv);dv.add(sheet[selector])
    for row in scenario['rows']:
        for address,value in row['cells'].items():
            sheet[address]=value;sheet[address].font=Font(name='Arial',color='0000FF')
    columns=list(scenario['table_columns']);start=13;end=13+len(scenario['rows'])-1;lookup_index=0
    for m in scenario['formula_metrics']:
        refs=m['required_references']
        if refs==[selector]:
            lookup_index+=1;column=columns[lookup_index]
            formula=f'=INDEX(${column}${start}:${column}${end},MATCH(${selector[0]}${selector[1:]},${columns[0]}${start}:${columns[0]}${end},0))'
        else:formula='='+'*'.join(refs)
        sheet[m['cell']]=formula;sheet[m['cell']].font=Font(name='Arial',color='000000')
    for i,spec in enumerate(req['charts']):
        origin=book[spec['required_references'][0].split('!')[0]]
        chart=LineChart();chart.title=spec['description']
        chart.add_data(Reference(origin,min_col=2,max_col=min(origin.max_column,4),min_row=1,max_row=min(origin.max_row,25)),titles_from_data=True)
        chart.set_categories(Reference(origin,min_col=1,min_row=2,max_row=min(origin.max_row,25)))
        sheet.add_chart(chart,'H'+str(5+i*18))
    for address in ('B5:B6','B15:B16'):
        sheet.conditional_formatting.add(address,CellIsRule(operator='greaterThan',formula=['0'],
            fill=PatternFill(fill_type='solid',start_color='D9E1F2',end_color='D9E1F2')))
    sheet['B5'].comment=Comment('Frozen source metric; cross-sheet reference.','Qualification')
    sheet[selector].comment=Comment('Illustrative scenario assumptions; independent fixture.','Qualification')
    result=output/manifest['output_workbook'];book.save(result);book.close();del book
    for item in manifest['publish_files']:shutil.copy2(workspace/'input'/item['input'],output/item['output'])
    with (output/'formula_recalc.json').open('w') as out:
        code=subprocess.run([sys.executable,str(ROOT/'vendor/skills/xlsx/scripts/recalc.py'),str(result),str(req['recalc_timeout'])],stdout=out,check=False).returncode
    if code:return code
    write_json(output/'xlsx_enhancement_summary.json',dict(execution='independent-fixture-qualification; not Agent-generated',
                helper_filename=None,recalc_invocation_count=1))
    report=verify(path,workspace,output/'business_verification.json')
    write_json(workspace.parent/'qualification.json',dict(status=report['status'],failures=report['failures'],
                execution='independent-fixture-qualification; not an Agent trace',dataset=args.dataset))
    print('qualification',report['status'],report['failures'],flush=True)
    return 0 if report['status']=='success' else 1


if __name__=='__main__':raise SystemExit(main())
