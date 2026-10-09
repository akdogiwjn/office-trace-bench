"""Independent calculations for the frozen native Census releases; never a task recipe."""
import re
from pathlib import Path
import sys
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import office
from openpyxl import load_workbook


def adjusted_retail_row(book, year):
    rows = list(book[str(year)].iter_rows(values_only=True))
    header = next(row for row in rows if any(isinstance(v, str) and v.startswith('Jan.') for v in row))
    adjusted = False
    for row in rows:
        if any(isinstance(v, str) and v.startswith('ADJUSTED(') for v in row):
            adjusted = True
        if adjusted and 'Retail and food services sales, total' in row:
            return {re.sub(r'\(p\)|\(r\)', '', str(label)).strip(): value
                    for label, value in zip(header, row) if isinstance(label, str) and re.match(r'^[A-Z][a-z]{2}\.?\s', label)}
    raise ValueError('missing adjusted retail total for ' + str(year))


def calculate(dataset, workbook):
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        if dataset == 'retail':
            current, prior = adjusted_retail_row(book, 2026), adjusted_retail_row(book, 2025)
            sales = current['Jul. 2026']
            values = dict(adjusted_sales=sales, monthly_growth=sales/current['Jun. 2026']-1,
                          annual_growth=sales/prior['Jul. 2025']-1)
            charts = dict(adjusted_trend=[current[m+' 2026'] for m in ('May', 'Jun.', 'Jul.')])
            assumptions = [('Base', 1), ('Higher', 1.04), ('Lower', .96)]
            scenarios = {name: dict(projected_sales=sales*factor) for name, factor in assumptions}
        elif dataset == 'manufacturing':
            rows = list(book['Table 1'].iter_rows(values_only=True))
            if not any(row[0] == 'January 2025 - March 2025' for row in rows):
                raise ValueError('unexpected M3 publication period')
            def record(prefix):
                found = [row for row in rows if isinstance(row[0], str) and row[0].strip().startswith(prefix)]
                if len(found) != 1:
                    raise ValueError('ambiguous M3 aggregate: ' + prefix)
                return found[0]
            total, durable, nondurable = [record(label) for label in
                ('All manufacturing industries', 'Durable goods industries', 'Nondurable goods industries')]
            # C/D/E are the published seasonally adjusted March/February/January
            # columns of this frozen release, independently inspected against its headers.
            values = dict(total_shipments=total[2], monthly_growth=total[2]/total[3]-1,
                          durable_share=durable[2]/total[2], nondurable_shipments=nondurable[2])
            charts = dict(durability_mix=[durable[2], nondurable[2]], shipment_trend=[total[4], total[3], total[2]])
            scenarios = {name: dict(projected_shipments=total[2]*factor) for name, factor in
                         [('Base', 1), ('Expansion', 1.03), ('Contraction', .97)]}
        else:
            raise ValueError('not a native Census dataset: ' + dataset)
        return dict(schema_version='office-semantic-oracle-v2', metric_values=values,
                    scenario_values=scenarios, chart_values=charts,
                    oracle_provenance='Independent calculations over the full unmodified native Census workbook. These values and source locations are not Agent task instructions.')
    finally:
        book.close()
