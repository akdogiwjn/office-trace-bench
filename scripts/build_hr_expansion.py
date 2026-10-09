#!/usr/bin/env python3
"""Freeze an OEWS-derived workforce fixture without inventing worker records."""
from collections import defaultdict
import csv
import gc
import hashlib
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import office
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, PieChart, Reference
from office_trace_bench.contracts import read_json, sha256, write_json
from office_trace_bench.workbooks import workbook_snapshot
from office_trace_bench.prompts import render_prompt
from build_xlsx_expansion import table


def read_source():
    source = ROOT / 'sources/xlsx/national_M2025_dl.xlsx'
    book = load_workbook(source, read_only=True, data_only=True)
    iterator = iter(book['national_M2025_dl'].values)
    headers = next(iterator)
    selected = []
    for values in iterator:
        row = dict(zip(headers, values))
        if (str(row['AREA']) == '99' and str(row['NAICS']) == '000000'
                and str(row['OWN_CODE']) == '1235' and row['O_GROUP'] == 'detailed'
                and all(isinstance(row[k], (int, float)) and row[k] > 0
                        for k in ('TOT_EMP', 'A_MEAN', 'A_MEDIAN'))):
            selected.append({k: row[k] for k in ('OCC_CODE', 'OCC_TITLE', 'TOT_EMP', 'A_MEAN', 'A_MEDIAN')})
    book.close()
    assert len(selected) > 500 and len({r['OCC_CODE'] for r in selected}) == len(selected)
    return source, sorted(selected, key=lambda r: r['OCC_CODE'])


def build():
    target = ROOT / 'datasets/xlsx/hr'
    if (target / 'manifest.json').exists():
        raise ValueError('HR dataset exists; refusing to overwrite frozen inputs')
    source, seeds = read_source()
    folder = target / 'input'; folder.mkdir(parents=True, exist_ok=True)
    book = Workbook(); raw = book.active; raw.title = 'Raw_Data'
    raw.append(['RecordID', 'SourceRowID', 'Dataset', 'Year', 'SOC', 'Occupation', 'SyntheticSegment',
                'EmploymentShard', 'AnnualMeanUSD', 'AnnualMedianUSD', 'PayrollMassUSD',
                'MedianWageMassUSD', 'SourceEmployment', 'ShardIndex', 'Origin', 'RowChecksum'])
    segments = defaultdict(lambda: [0, 0, 0])
    total_emp = payroll = median_mass = 0
    for index, seed in enumerate(seeds):
        employment = int(seed['TOT_EMP']); mean = int(seed['A_MEAN']); median = int(seed['A_MEDIAN'])
        total_emp += employment; payroll += employment * mean; median_mass += employment * median
        repeats = 100000 // len(seeds) + (index < 100000 % len(seeds))
        shard_sum = 0
        for shard in range(repeats):
            count = employment // repeats + (shard < employment % repeats)
            shard_sum += count
            segment = 'Benchmark segment ' + str(shard % 6 + 1)
            amounts = [count, count * mean, count * median]
            segments[segment] = [a + b for a, b in zip(segments[segment], amounts)]
            identifier = f'hr-{index + 1:04d}-{shard + 1:05d}'
            raw.append([identifier, f'source-{index + 1:04d}', 'hr', 2025, seed['OCC_CODE'],
                        seed['OCC_TITLE'], segment, count, mean, median, count * mean,
                        count * median, employment, shard + 1, 'synthetic_partition',
                        hashlib.sha256(identifier.encode()).hexdigest()[:24]])
        assert shard_sum == employment
    assert raw.max_row == 100001 and raw.max_column == 16
    assert sum(v[0] for v in segments.values()) == total_emp
    assert sum(v[1] for v in segments.values()) == payroll
    raw.freeze_panes = 'A2'; raw.auto_filter.ref = raw.dimensions
    occupation = table(book, 'Occupation_Summary', ['SOC', 'Occupation', 'Employment', 'Annual Mean USD',
        'Annual Median USD', 'Estimated annual payroll mass USD'],
        [[s['OCC_CODE'], s['OCC_TITLE'], s['TOT_EMP'], s['A_MEAN'], s['A_MEDIAN'],
          int(s['TOT_EMP']) * int(s['A_MEAN'])] for s in seeds])
    largest = max(seeds, key=lambda s: s['TOT_EMP'])['OCC_TITLE']
    occupation['H1'] = 'Highest employment selected occupation'; occupation['H2'] = largest
    table(book, 'Employment_Summary', ['Metric', 'Value'], [['Selected employment', total_emp],
        ['Estimated payroll mass USD', payroll], ['Occupation median wage mass USD', median_mass]])
    table(book, 'Wage_Summary', ['Metric', 'Value'], [
        ['Employment-weighted annual mean wage', '=Employment_Summary!B3/Employment_Summary!B2'],
        ['Employment-weighted occupation median proxy', '=Employment_Summary!B4/Employment_Summary!B2']])
    table(book, 'Segment_Summary', ['Synthetic segment', 'Employment', 'Payroll mass USD', 'Median wage mass USD'],
        [[name, *amounts] for name, amounts in sorted(segments.items())])
    top = sorted(seeds, key=lambda s: (-s['TOT_EMP'], s['OCC_CODE']))[:24]
    mix = table(book, 'Employment_Mix', ['Occupation', 'Employment', 'Annual mean USD', 'Annual median USD'],
        [[s['OCC_TITLE'], s['TOT_EMP'], s['A_MEAN'], s['A_MEDIAN']] for s in top])
    recon = table(book, 'Reconciliation', ['Check', 'Value'], [
        ['Raw selected employment', '=SUM(Raw_Data!H2:H100001)'], ['Source selected employment', total_emp],
        ['Employment agrees', '=IF(B2=B3,"PASS","REVIEW")'], ['Raw payroll mass', '=SUM(Raw_Data!K2:K100001)'],
        ['Source payroll mass', payroll], ['Payroll agrees', '=IF(B5=B6,"PASS","REVIEW")'],
        ['Raw median wage mass', '=SUM(Raw_Data!L2:L100001)'], ['Source median wage mass', median_mass],
        ['Median mass agrees', '=IF(B8=B9,"PASS","REVIEW")'], ['Raw records', '=COUNTA(Raw_Data!A2:A100001)'],
        ['Expected records', 100000], ['Row count agrees', '=IF(B11=B12,"PASS","REVIEW")']])
    recon['A15'] = 'Overall status'; recon['B15'] = '=IF(AND(B4="PASS",B7="PASS",B10="PASS",B13="PASS"),"PASS","REVIEW")'
    line = LineChart(); line.title = 'Existing selected wages'
    line.add_data(Reference(mix, min_col=3, max_col=4, min_row=1, max_row=25), titles_from_data=True)
    line.set_categories(Reference(mix, min_col=1, min_row=2, max_row=25)); mix.add_chart(line, 'F2')
    pie = PieChart(); pie.title = 'Existing synthetic segment mix'; segment_sheet = book['Segment_Summary']
    pie.add_data(Reference(segment_sheet, min_col=2, min_row=1, max_row=7), titles_from_data=True)
    pie.set_categories(Reference(segment_sheet, min_col=1, min_row=2, max_row=7)); segment_sheet.add_chart(pie, 'F2')
    baseline = workbook_snapshot(book, ['Raw_Data'])
    formula_count = sum(len(v) for v in baseline['formulas'].values())
    workbook = folder / 'workforce_template.xlsx'; book.save(workbook); book.close(); del book; gc.collect()
    prepared = []
    for name, header, rows in [
        ('prepared_workforce_summary.csv', ['selected_employment', 'weighted_mean_wage_usd', 'weighted_occupation_median_proxy_usd'],
         [[total_emp, payroll / total_emp, median_mass / total_emp]]),
        ('prepared_reconciliation_summary.csv', ['raw_rows', 'selected_source_occupations', 'partition_totals_preserved'],
         [[100000, len(seeds), True]])]:
        path = folder / name
        with path.open('w', newline='') as stream:
            writer = csv.writer(stream); writer.writerow(header); writer.writerows(rows)
        prepared.append(path)
    url = 'https://raw.githubusercontent.com/hack4rva/richmond-ai-impact-analysis/main/data/oesm25nat/national_M2025_dl.xlsx'
    provenance = dict(publisher='U.S. Bureau of Labor Statistics', distributor='Public GitHub repository cache',
        official_source_url='https://www.bls.gov/oes/special-requests/oesm25nat.zip', source_url=url,
        source_artifact=str(source.relative_to(ROOT)), source_sha256=sha256(source), year=2025,
        official_download_available=False, official_byte_identity_independently_verified=False,
        observed_seed_rows=len(seeds), benchmark_raw_rows=100000, raw_rows_are_observed_workers=False,
        selection='U.S. cross-industry ownership 1235 detailed occupations with numeric positive employment, annual mean and annual median; exclude all aggregate hierarchy rows and suppressed wages.',
        expansion='Partition employment counts into synthetic shards; preserve source occupation counts and count-times-wage masses.',
        median_definition='Employment-weighted average of occupation medians; a proxy, never the national worker median.',
        payroll_definition='Employment times published annual mean wage; illustrative estimated wage mass, not observed payroll.',
        industries='No industry detail in national cross-industry source; segments are explicitly synthetic.')
    write_json(folder / 'template_manifest.json', dict(dataset_id='hr', raw_count=100000,
        base_sheets=baseline['sheet_order'], existing_formula_count=formula_count, existing_chart_count=2, source=provenance))
    (folder / 'dataset_sources.md').write_text('# BLS OEWS workforce benchmark\n\nSource publisher: BLS; 2025 national workbook obtained through a public repository cache because official download returned 403.\n'
        'Cache URL: ' + url + '\nOfficial byte identity is not independently verified. Freeze by recorded SHA-256.\n'
        + provenance['selection'] + '\n' + provenance['expansion'] + '\n'
        '100,000 synthetic workload shards are not observed worker records; seven base sheets and two existing charts.\n'
        + provenance['median_definition'] + '\n' + provenance['payroll_definition'] + '\n' + provenance['industries'] + '\n'
        'Do not sum hierarchy totals together with detailed occupations. Coverage is the selected occupations only.\n'
        'Preserve all raw values, formulas and charts. Read dataset_manifest.json for KPI/scenario requirements.\n')
    metrics = [dict(name=name, cell=f'B{i}', required_references=[ref], font_color='008000',
                    **({'number_format': fmt} if fmt else {}))
        for i, (name, ref, fmt) in enumerate([
            ('Selected Employment', 'Employment_Summary!B2', '#,##0'),
            ('Employment-weighted Annual Mean Wage', 'Wage_Summary!B2', '$#,##0.00'),
            ('Employment-weighted Occupation Median Proxy', 'Wage_Summary!B3', '$#,##0.00'),
            ('Highest Employment Selected Occupation', 'Occupation_Summary!H2', None),
            ('Reconciliation Status', 'Reconciliation!B15', None)], 5)]
    scenario = dict(selector_cell='B12', default='Base', table_columns={'D':'Scenario', 'E':'Wage multiplier', 'F':'Headcount multiplier'},
        rows=[dict(name=name, cells={f'D{i}':name, f'E{i}':wage, f'F{i}':count})
              for i, (name, wage, count) in enumerate([('Base', 1., 1.), ('Upside', 1.08, 1.05), ('Downside', .93, .95)], 13)],
        formula_metrics=[dict(name=name, cell=cell, required_references=refs, meaning=meaning, font_color='000000')
          for name, cell, refs, meaning in [
            ('Selected wage multiplier', 'B13', ['B12'], 'Look up selected wage multiplier'),
            ('Selected headcount multiplier', 'B14', ['B12'], 'Look up selected headcount multiplier'),
            ('Projected Employment', 'B15', ['B5', 'B14'], 'Multiply selected employment by headcount multiplier'),
            ('Projected Estimated Payroll', 'B16', ['B5', 'B6', 'B13', 'B14'], 'Multiply employment by annual mean wage and both scenario multipliers')]])
    write_json(target / 'expected.json', dict(baseline=baseline,
        cached_values=dict(B5=total_emp, B6=payroll/total_emp, B7=median_mass/total_emp, B8=largest,
                           B9='PASS', B13=1, B14=1, B15=total_emp, B16=payroll),
        oracle_provenance='Independent pre-Agent aggregation of selected detailed occupations; employment and wage-mass conservation.'))
    manifest = dict(schema_version='office-dataset-v1', kind='xlsx', dataset_id='hr', domain='workforce',
        task='Enhance the BLS OEWS selected workforce benchmark workbook', workbook=workbook.name,
        output_workbook='workforce_report.xlsx', requirements=dict(summary_sheet='Executive_Summary',
            base_sheets=baseline['sheet_order'], raw_sheets=['Raw_Data'], minimum_bytes=1000000,
            minimum_formula_count=formula_count+9, metrics=metrics, scenario=scenario,
            charts=[dict(id='employment_mix', description='employment mix of selected occupations', required_references=['Employment_Mix!']),
                    dict(id='wage_distribution', description='annual mean and median wage comparison of selected occupations', required_references=['Occupation_Summary!'])],
            minimum_conditional_formats=2, minimum_comments=2, recalc_timeout=180),
        publish_files=[dict(input=p.name, output=p.name.removeprefix('prepared_'), sha256=sha256(p)) for p in prepared],
        required_outputs=['workforce_report.xlsx', 'workforce_summary.csv', 'reconciliation_summary.csv',
                          'formula_recalc.json', 'business_verification.json', 'xlsx_enhancement_summary.json'],
        input_files={str(p.relative_to(target)):sha256(p) for p in sorted(folder.iterdir()) if p.is_file()},
        expected_file='expected.json', expected_sha256=sha256(target/'expected.json'), prompt_template='xlsx.txt',
        prompt_template_sha256=sha256(ROOT/'prompts/xlsx.txt'),
        prompt_contract=dict(task_subject='BLS OEWS 2025 selected detailed-occupation workforce benchmark workbook',
            raw_record_count=100000, original_chart_count=2, metadata_files=['template_manifest.json', 'dataset_sources.md']),
        runtime_case='SUB-MEM-OFFICE-HR-01', agent_context=read_json(ROOT/'datasets/xlsx/tlc/manifest.json')['agent_context'],
        trace_rules=read_json(ROOT/'datasets/xlsx/retail/manifest.json')['trace_rules'], source_provenance=provenance)
    write_json(target / 'manifest.json', manifest)
    render_prompt(manifest, '/validation/workspace', target/'manifest.json')
    print('hr', len(seeds), 'selected occupations; 100000 synthetic rows;', formula_count, 'existing formulas;', workbook.stat().st_size, 'bytes', flush=True)


if __name__ == '__main__':
    build()
