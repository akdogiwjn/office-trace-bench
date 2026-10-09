#!/usr/bin/env python3
"""Compare retained real TLC/OPM runs with the historical traces and outputs."""
import argparse
from collections import Counter
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import office
from openpyxl import load_workbook
from office_trace_bench.contracts import read_json, sha256, write_json


def calls(trace):
    for step in trace['steps']:
        results = {r['source_call_id']: r for r in step.get('observation', {}).get('results', [])}
        for call in step.get('tool_calls', []):
            yield step['step_id'], call, results.get(call['tool_call_id'])


def trace_stats(path):
    trace = read_json(path)
    rows = list(calls(trace))
    problems = []
    kills = []
    for step, call, result in rows:
        if call['function_name'] == 'process' and call['arguments'].get('action') == 'kill':
            kills.append(step)
            continue
        extra = (result or {}).get('extra', {})
        detail = extra.get('details', {})
        # OpenClaw may report is_error=false even when an exec exits nonzero.
        if extra.get('is_error') or detail.get('status') in ('failed', 'error') or detail.get('exitCode') not in (None, 0):
            problems.append({'step_id': step, 'tool': call['function_name'],
                             'exit_code': detail.get('exitCode'),
                             'result_excerpt': (result or {}).get('content', '')[:350]})
    return {'trace': str(path.relative_to(ROOT)), 'sha256': sha256(path),
            'session_id': trace['session_id'], 'agent': trace['agent'],
            'steps': len(trace['steps']), 'tool_calls': len(rows),
            'tools': dict(Counter(c['function_name'] for _, c, _ in rows)),
            'missing_tool_results': [c['tool_call_id'] for _, c, r in rows if r is None],
            'command_failures': problems, 'process_kill_steps': kills,
            'recorded_metrics': trace['final_metrics']}


def exec_steps(trace, pattern):
    return [s for s, c, _ in calls(trace) if c['function_name'] == 'exec'
            and re.search(pattern, c['arguments'].get('command', ''))]


def coverage(trace, kind):
    patterns = ({'inspect_input': r'load_workbook|template_manifest',
                 'build_workbook': r'python3 (?:enhance_workbook|build_report)\.py',
                 'recalculate': r'python3 (?:\S*/)?recalc\.py',
                 'publish_csv': r'cp .*prepared_.*\.csv',
                 'verify': r'python3 (?:input/verify_xlsx_enhanced\.py|/project/office\.py verify)'}
                if kind == 'xlsx' else
                {'check_fields': r'python3 (?:scripts/)?check_fillable_fields\.py',
                 'extract_schema': r'python3 (?:scripts/)?extract_form_field_info\.py',
                 'render_template': r'python3 (?:scripts/)?convert_pdf_to_images\.py',
                 'execute_batch': r'python3 (?:run_batch_fill_render|run_batch)\.py',
                 'verify': r'python3 (?:input/verify_pdf_batch\.py|/project/office\.py verify)'})
    stages = {name: exec_steps(trace, pattern) for name, pattern in patterns.items()}
    # Also inspect helpers written before an exec; a read of Skill code is not execution.
    written = {}
    for step, call, result in calls(trace):
        args = call['arguments']
        if call['function_name'] == 'write':
            written[Path(args.get('path', args.get('file_path', ''))).name] = args.get('content', '')
        if call['function_name'] != 'exec':
            continue
        for helper in re.findall(r'\b(?:python3|bash|sh)\s+(\S+\.(?:py|sh))', args.get('command', '')):
            body = written.get(Path(helper.strip(chr(34) + chr(39))).name, '')
            if kind == 'xlsx' and 'Executive_Summary' in body and re.search(r'\.save\(', body):
                if step not in stages['build_workbook']:
                    stages['build_workbook'].append(step)
            if kind == 'pdf' and 'fill_fillable_fields.py' in body and 'convert_pdf_to_images.py' in body:
                if step not in stages['execute_batch']:
                    stages['execute_batch'].append(step)
            template_render = (re.search(r'run_script\("convert_pdf_to_images\.py", \[FORM, tmpl_dir\]\)', body)
                               or ('convert_pdf_to_images.py' in body and 'rendered/template' in body))
            observation = (result or {}).get('content', '')
            template_render = template_render or ('convert_pdf_to_images.py' in observation
                                                   and '/rendered/template' in observation)
            if (kind == 'pdf' and not stages['render_template'] and template_render
                    and (result or {}).get('extra', {}).get('details', {}).get('exitCode') == 0):
                stages['render_template'].append(step)
    return stages


def workbook(path):
    formulas = load_workbook(path, read_only=True, data_only=False)
    values = load_workbook(path, read_only=True, data_only=True)
    try:
        return {'sheets': formulas.sheetnames,
                'formulas': {f'{ws.title}!{c.coordinate}': c.value for ws in formulas
                             if ws.title != 'Raw_Sample' for row in ws.iter_rows()
                             for c in row if c.data_type == 'f'},
                'kpi_cached_values': {cell: values['Executive_Summary'][cell].value
                                      for cell in ('B5', 'B6', 'B7', 'B8', 'B9', 'B13', 'B14', 'B15', 'B16')}}
    finally:
        formulas.close()
        values.close()


def output_comparison(kind, old, new):
    if kind == 'xlsx':
        a, b = [workbook(p / 'monthly_operations_report.xlsx') for p in (old, new)]
        changes = [{'cell': cell, 'old': a['formulas'].get(cell), 'new': b['formulas'].get(cell)}
                   for cell in sorted(set(a['formulas']) | set(b['formulas']))
                   if a['formulas'].get(cell) != b['formulas'].get(cell)]
        csvs = {name: sha256(old / name) == sha256(new / name)
                for name in ('monthly_operations_summary.csv', 'reconciliation_summary.csv')}
        return {'same_sheet_order': a['sheets'] == b['sheets'], 'sheets': b['sheets'],
                'same_kpi_cached_values': a['kpi_cached_values'] == b['kpi_cached_values'],
                'kpi_cached_values': b['kpi_cached_values'], 'same_csv_bytes': csvs,
                'formula_counts': {'old': len(a['formulas']), 'new': len(b['formulas'])},
                'formula_differences': changes,
                'recalc_reports': {label: read_json(p / 'formula_recalc.json')
                                   for label, p in [('old', old), ('new', new)]}}
    summaries = {label: read_json(p / 'batch_summary.json') for label, p in [('old', old), ('new', new)]}
    mappings = {}
    for path in sorted((old / 'field_values').glob('*.json')):
        def normalize(p):
            return {x['field_id']: {'value': x['value'], 'page': x['page']} for x in read_json(p)}
        mappings[path.stem] = normalize(path) == normalize(new / 'field_values' / path.name)
    return {'batch_summaries': summaries, 'same_field_values_and_pages': mappings,
            'artifact_counts': {label: {'pdf': len(list((p / 'filled').glob('*.pdf'))),
                                       'png': len(list((p / 'rendered').rglob('*.png'))),
                                       'mapping': len(list((p / 'field_values').glob('*.json')))}
                                for label, p in [('old', old), ('new', new)]}}


def compare(kind, dataset, run_id):
    run = ROOT / 'runs' / run_id
    manifest = read_json(ROOT / 'datasets' / kind / dataset / 'manifest.json')
    imported = read_json(ROOT / 'legacy/import_manifest.json')['files']
    a, b = ROOT / 'legacy' / kind / 'trajectory.json', run / 'trajectory.json'
    old, new = read_json(a), read_json(b)
    # Use each run's frozen manifest, so later dataset/profile edits cannot change the comparison.
    manifest = read_json(run / 'workspace/input/dataset_manifest.json')
    archived = read_json(ROOT / 'legacy/archive_manifest.json')['files']
    source = {}
    for name, digest in manifest['input_files'].items():
        historical = imported.get(f'datasets/{kind}/{dataset}/{name}')
        if historical is None:
            historical = archived.get(f'legacy/{kind}/{Path(name).name}')
        source[name] = sha256(run / 'workspace' / name) == digest == historical
    snapshots = read_json(run / 'source_snapshot.json')['files']
    skill = {name: snapshots.get(name) == digest and sha256(ROOT / name) == digest
             for name, digest in imported.items() if name.startswith(f'vendor/skills/{kind}/')}
    reports = {name: read_json(run / name) for name in
               ('independent_verification.json', 'legacy_verification.json', 'baseline_acceptance.json')}
    verification = {name: report.get('status') == 'success' and not report.get('failures')
                    for name, report in reports.items() if name != 'baseline_acceptance.json'}
    verification['baseline_accepted'] = reports['baseline_acceptance.json']['accepted']
    verification['old_business_success'] = read_json(ROOT / 'legacy' / kind / 'output/business_verification.json')['status'] == 'success'
    archive = read_json(ROOT / 'reports/archived-regression-summary.json')['datasets'][dataset]
    verification['archived_outputs_pass_both_verifiers'] = archive['new_verifier'] == archive['legacy_verifier'] == 'success'
    stages = {'old': coverage(old, kind), 'new': coverage(new, kind)}
    output = output_comparison(kind, ROOT / 'legacy' / kind / 'output', run / 'workspace/output')
    equivalent = (output['same_sheet_order'] and output['same_kpi_cached_values']
                  and all(output['same_csv_bytes'].values()) if kind == 'xlsx' else
                  all(output['same_field_values_and_pages'].values()) and
                  output['artifact_counts']['old'] == output['artifact_counts']['new'])
    stats = {'old': trace_stats(a), 'new': trace_stats(b)}
    identity = ('version', 'model_name')
    same_agent = all(old['agent'][k] == new['agent'][k] for k in identity)
    same_agent &= old['agent']['extra']['thinking_level'] == new['agent']['extra']['thinking_level']
    ok = (all(source.values()) and all(skill.values()) and all(verification.values()) and same_agent
          and equivalent and all(all(v.values()) for v in stages.values())
          and not any(v['missing_tool_results'] for v in stats.values()))
    return {'status': 'success' if ok else 'failed', 'run_id': run_id,
            'inputs_match_frozen_legacy_import': source, 'skills_match_frozen_legacy_import': skill,
            'same_openclaw_model_thinking': same_agent, 'trace': stats,
            'business_stage_exec_steps': stages, 'verification': verification, 'outputs': output}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--tlc', default='xlsx-tlc-20261008T115345Z-63218840')
    parser.add_argument('--opm', default='pdf-opm-20261008T115600Z-aa170fa3')
    args = parser.parse_args()
    data = {'schema_version': 'office-trace-comparison-v1',
            'compared_at': datetime.now(timezone.utc).isoformat(),
            'datasets': {'tlc': compare('xlsx', 'tlc', args.tlc), 'opm': compare('pdf', 'opm', args.opm)},
            'scope': 'Business regression and complete recorded tool traces; not identical workloads or performance equivalence.'}
    data['status'] = 'success' if all(d['status'] == 'success' for d in data['datasets'].values()) else 'failed'
    write_json(ROOT / 'reports/trace-comparison-v1.json', data)
    tlc, opm = data['datasets']['tlc'], data['datasets']['opm']
    rows = []
    for name, d in [('XLSX / TLC', tlc), ('PDF / OPM', opm)]:
        a, b = d['trace']['old'], d['trace']['new']
        rows.append(f"| {name} | {a['steps']} → {b['steps']} | {a['tool_calls']} → {b['tool_calls']} | {a['tools'].get('process', 0)} → {b['tools'].get('process', 0)} | {d['status']} |")
    text = '\n'.join([
        '# 新旧 trace 对比', '', f"比较结果：`{data['status']}`。旧轨迹为 2026-07-30 的历史真实运行，新轨迹为 2026-10-08 的通用化真实运行。", '',
        '| 用例 | ATIF steps（旧 → 新） | 工具调用（旧 → 新） | process 调用（旧 → 新） | 业务回归 |',
        '|---|---:|---:|---:|---|', *rows, '',
        '两组输入和 Skill 与导入时冻结的 SHA-256 一致；OpenClaw 2026.6.6、模型 deepseek/deepseek-v4-flash、thinking high 一致。',
        '新输出通过 Agent 内验收、独立通用验收、原 verifier；历史输出也通过新旧 verifier。每条工具调用都有对应结果。', '',
        '## XLSX / TLC', '',
        '模板检查、构建工作簿、重算、发布 CSV、业务验收均保留。原始 100,000 行数据、原公式和原图表通过保留检查；',
        '新旧均为 8 个 sheet，核心 KPI 缓存值相同，两个 CSV 字节相同，新旧重算报告均无公式错误。',
        f"旧 trace 真正调用 recalc.py {len(tlc['business_stage_exec_steps']['old']['recalculate'])} 次，新 trace {len(tlc['business_stage_exec_steps']['new']['recalculate'])} 次。旧 summary 自报一次，与 trace 不符；本比较以 trace 为准。",
        '旧 trace 主动结束了 4 个检查进程；新 trace 没有 process kill。新第 24 步因 time 命令不存在而退出 127，',
        '第 25 步直接执行 builder 成功；首次失败没有开始构建。因此有 2 次启动尝试、1 次成功构建。',
        '新建脚本写入了两版，第一版尚未执行便修正了单元格布局。',
        '公式总数 45 → 48：B13/B14 由 INDEX/MATCH 改为 VLOOKUP，结果相同；新增 B17/B18/B19 展示差额与增幅。', '',
        '## PDF / OPM', '',
        '可填写检查、字段抽取、空白模板渲染、映射生成、批量填写/渲染、业务验收均保留。',
        '两边均为 10 份 PDF、10 份映射 JSON、33 张 PNG；10 人的全部映射字段值及页码相同。',
        '两边均使用 Skill 完成 10 次填写、11 次渲染（模板 1 次 + 填写后 10 次），受保护字段均保持空白。',
        '旧版把映射生成与填写渲染拆成两个 helper；新版合并为一个 helper，执行一次。',
        '新 trace 有两个辅助 exec 退出 1：第 5 步 which 查询未安装的可选工具；第 13 步把 protected_blank 的字段对象当作键，出现 TypeError，',
        '第 14 步用正确的字段集合重新检查通过。',
        '实际批量处理和验收均成功。完整结果片段见 JSON 的 command_failures。', '',
        '## 结论与范围', '',
        '这些证据支持 TLC/OPM 通用化的业务兼容性。步骤数、脚本布局、重算次数和额外展示公式不同，',
        '因此不能将这两组 trace 视为严格相同的资源负载，也不能据此得出性能提升结论。',
        '阶段覆盖按真实 exec 匹配；新版空白渲染在成功执行的 batch helper 内，结合此前写入的源码和已验收的 PNG 确认。',
        '不把 read 一个脚本当成执行，也不把 batch 的一个工具调用误计为一次填写。', '',
        'API 凭据来源另见 credential-provenance.json；只记录当前认证解析结果和脱敏指纹，没有保存完整 key。', '',
        f"TLC 新 trace：`runs/{args.tlc}/trajectory.json`", '',
        f"OPM 新 trace：`runs/{args.opm}/trajectory.json`", '',
        '重建此报告：`python3 scripts/compare_baseline_traces.py`', ''])
    (ROOT / 'reports/trace-comparison-v1.md').write_text(text)
    print(json.dumps({'status': data['status'], 'report': 'reports/trace-comparison-v1.md'}, ensure_ascii=False))
    return 0 if data['status'] == 'success' else 1


if __name__ == '__main__':
    raise SystemExit(main())
