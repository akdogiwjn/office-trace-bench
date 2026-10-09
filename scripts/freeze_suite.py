#!/usr/bin/env python3
"""Explicit v2 migration; source observations are never expanded or sharded."""
import argparse
from datetime import datetime, timezone
from pathlib import Path
import re
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import office
from openpyxl import Workbook, load_workbook
from office_trace_bench.complexity import inspect_workbook
from office_trace_bench.contracts import read_json, write_json, sha256


def lines(path, page=None):
    result = {}
    for line in path.read_text().splitlines():
        match = re.match(r'L(\d+)@P([^:]+): (.*)', line)
        if match and (page is None or match[2] == str(page)):
            result[int(match[1])] = match[3]
    return result


def table_rows(source, start, stop, count):
    rows = []; pending = ''
    token = r'(?:-?\d[\d,]*(?:\.\d+)?|\((?:\*|S|NA)\))'
    ending = re.compile(r'(?P<data>' + token + r'(?:\s+' + token + r'){' + str(count-1) + r'})\s*$')
    for i in range(start, stop + 1):
        text = source[i]
        match = ending.search(text)
        if not match:
            if text.strip() not in ('2',): pending += ' ' + text.strip()
            continue
        label = pending + ' ' + text[:match.start()]
        label = re.sub(r'[.…]+', ' ', label)
        label = re.sub(r'\s+', ' ', label).strip()
        pending = ''
        values = []
        for value in match['data'].split():
            if value.startswith('('): values.append(value)
            else: values.append(float(value.replace(',', '')) if '.' in value else int(value.replace(',', '')))
        rows.append([label, *values])
    if pending.strip(): raise ValueError('unparsed table text: ' + pending)
    return rows


def metric(id, name, meaning, unit):
    return dict(id=id, name=name, meaning=meaning, unit=unit)

def chart_oracle(dataset, workbook):
    """Independent source-value oracle; never included in the task manifest."""
    book=load_workbook(workbook,read_only=True,data_only=True)
    try:
        if dataset=='tlc':
            return dict(daily_trips=[r[0] for r in book['Daily_Summary'].iter_rows(min_row=2,max_row=32,min_col=3,max_col=3,values_only=True)],
                        payment_mix=[r[0] for r in book['Payment_Summary'].iter_rows(min_row=2,max_row=5,min_col=2,max_col=2,values_only=True)])
        rows=list(book.worksheets[0].iter_rows(values_only=True))
        if dataset=='retail':return dict(adjusted_trend=[rows[1][10],rows[1][9],rows[1][8]])
        if dataset=='manufacturing':
            durable=next(r for r in rows if r[0]=='Durable goods industries')
            nondurable=next(r for r in rows if r[0]=='Nondurable goods industries')
            return dict(durability_mix=[durable[1],nondurable[1]],shipment_trend=[rows[1][3],rows[1][2],rows[1][1]])
        if dataset=='hr':
            records=[dict(zip(rows[0],r)) for r in rows[1:]]
            major=[r for r in records if r['O_GROUP']=='major' and r['NAICS']=='000000' and str(r['OWN_CODE'])=='1235']
            return dict(major_employment=[r['TOT_EMP'] for r in sorted(major,key=lambda r:r['TOT_EMP'],reverse=True)[:5]])
        raise ValueError('unknown migration input')
    finally:book.close()


def freeze(dataset, workbook, metrics, values, scenario, projections, chart_specs, provenance):
    target = ROOT / 'datasets/xlsx' / dataset
    target.mkdir(parents=True, exist_ok=True)
    if (target/'manifest.json').exists() and read_json(target/'manifest.json').get('schema_version') == 'office-dataset-v2':
        raise ValueError('v2 input is frozen; publish a new suite revision instead of replacing it')
    inputs = target/'input'; inputs.mkdir(exist_ok=True)
    if workbook.parent != inputs: shutil.copy2(workbook, inputs/workbook.name)
    workbook = inputs/workbook.name
    write_json(target/'provenance.json', dict(schema_version='office-input-provenance-v2', dataset_id=dataset,
        filename=workbook.name, sha256=sha256(workbook), frozen_at=datetime.now(timezone.utc).isoformat(), **provenance))
    write_json(target/'complexity.json', inspect_workbook(workbook))
    write_json(target/'expected.json', dict(schema_version='office-semantic-oracle-v2', metric_values=values,
        scenario_values=projections, chart_values=chart_oracle(dataset,workbook),
        oracle_provenance='Independent calculations over frozen source observations; these answers are not part of the Agent-facing task contract.'))
    context = read_json(ROOT/'datasets/pdf/opm/manifest.json')['agent_context']
    old = read_json(target/'manifest.json') if (target/'manifest.json').exists() else {}
    # Existing TLC verifier/layout remains in legacy; the formal Agent chooses its own implementation.
    inputs_info = {str(p.relative_to(target)): sha256(p) for p in inputs.iterdir() if p.is_file() and p.name != 'verify_xlsx_enhanced.py'}
    manifest = dict(schema_version='office-dataset-v2', workload_contract='spreadsheet-semantic-v2', kind='xlsx',
        dataset_id=dataset, domain=dataset, task=provenance['task'], workbook=workbook.name,
        output_workbook='monthly_operations_report.xlsx' if dataset=='tlc' else dataset+'_enhanced.xlsx',
        input_files=inputs_info, expected_file='expected.json', expected_sha256=sha256(target/'expected.json'),
        analysis_metadata={name:sha256(target/name) for name in ('provenance.json','complexity.json')},
        prompt_template='xlsx.txt', prompt_template_sha256=sha256(ROOT/'prompts/xlsx.txt'), agent_context=context,
        runtime_case=old.get('runtime_case', 'OFFICE-XLSX-'+dataset.upper()),
        requirements=dict(metrics=metrics, scenario=scenario, charts=chart_specs, minimum_comments=2,
                          minimum_conditional_formats=2, recalc_timeout=180),
        publish_files=old.get('publish_files', []),
        trace_rules=[dict(pattern=r'\bverify_office\.py\b',operation='verify_business'),
                     dict(pattern=r'recalc\.py',operation='recalculate_formulas')])
    manifest['required_outputs'] = [manifest['output_workbook'], 'workbook_features.json', 'formula_recalc.json',
        'business_verification.json', 'xlsx_enhancement_summary.json', *[p['output'] for p in manifest['publish_files']]]
    write_json(target/'manifest.json',manifest)


def build_tables():
    retail = lines(ROOT/'sources/provenance-audit/retail-m3-tables-web.txt',4)
    m3 = lines(ROOT/'sources/provenance-audit/m3-table1-web.txt',0)
    assert set(range(147,193)) <= set(retail)
    assert set(range(8,107)) <= set(m3)
    rows = table_rows(retail,147,192,12)
    book=Workbook();s=book.active;s.title='MARTS Table 1'
    s.append(['Kind of Business','NSA 2026 Jan-Aug','NSA YTD Change %','NSA Aug 2026','NSA Jul 2026','NSA Jun 2026',
              'NSA Aug 2025','NSA Jul 2025','SA Aug 2026','SA Jul 2026','SA Jun 2026','SA Aug 2025','SA Jul 2025'])
    for row in rows: s.append(row)
    notes=book.create_sheet('Publication notes')
    notes.append(['Advance monthly retail sales, August 2026; release CB26-153, September 16, 2026'])
    notes.append(['Frozen pre-September-28 revision, not current latest estimates. Values in USD millions.'])
    for i in range(193,208): notes.append([retail[i]])
    notes.append(['Official PDF table transcribed to XLSX; not native Census XLSX, no synthetic expansion.'])
    p=ROOT/'datasets/xlsx/retail/input/marts_table1_202608.xlsx';p.parent.mkdir(parents=True,exist_ok=True);book.save(p);book.close()
    total=rows[0];sa=total[8];prev=total[9];year=total[11]
    metric_values=dict(adjusted_sales=sa,monthly_growth=sa/prev-1,annual_growth=sa/year-1)
    assumptions=[dict(name=n,assumptions=dict(sales_multiplier=v)) for n,v in [('Base',1),('Higher',1.04),('Lower',.96)]]
    scenario=dict(default='Base',dimensions=[dict(id='sales_multiplier',meaning='Hypothetical sales level multiplier, not a Census forecast')],rows=assumptions,
                  projections=[dict(id='projected_sales',meaning='Adjusted sales scaled by the selected sales multiplier',unit='USD million')])
    freeze('retail',p,[metric('adjusted_sales','Adjusted retail and food services sales','Published total retail and food services seasonally adjusted sales for August 2026. Use the total, not a sum of overlapping categories.','USD million'),
        metric('monthly_growth','Month-over-month growth','Relative change in the adjusted total from July to August 2026; compute from sales levels rather than rounded published percentages.','fraction'),
        metric('annual_growth','Year-over-year growth','Relative change in the adjusted total from August 2025 to August 2026; compute from sales levels.','fraction')],metric_values,scenario,
        {r['name']:dict(projected_sales=sa*r['assumptions']['sales_multiplier']) for r in assumptions},
        [dict(id='adjusted_trend',meaning='June, July and August 2026 adjusted total sales; chronological order',unit='USD million')],
        dict(task='Analyze the frozen Census MARTS August 2026 publication table',input_class='official_table_transcription',native_xlsx=False,
             official_url='https://www.census.gov/retail/marts/www/marts_current.pdf',official_xlsx_url='https://www.census.gov/retail/marts/www/marts_current.xlsx',
             actual_download_source='Official census.gov PDF table via web-tool text extraction',source_filename='marts_current.pdf',
             release='CB26-153; August 2026 advance estimates, released 2026-09-16, before 2026-09-28 revisions',download_date='2026-10-09',
             source_evidence='sources/provenance-audit/retail-m3-tables-web.txt',source_evidence_sha256=sha256(ROOT/'sources/provenance-audit/retail-m3-tables-web.txt'),
             observations=len(rows),synthetic_expansion=False,conversion='All Table 1 rows and 12 numeric/marker columns transcribed; labels joined across PDF lines; publication notes retained. Source PDF and native XLSX byte hashes unavailable.',
             limitations=['Native Census XLSX downloads return HTTP 403; this fallback is explicitly generated from an official raw table, not an official XLSX file.']))
    rows=table_rows(m3,8,90,13)
    book=Workbook();s=book.active;s.title='Shipments August 2026'
    s.append(['Industry','SA Aug 2026','SA Jul 2026','SA Jun 2026','SA Jul-Aug Change %','SA Jun-Jul Change %','SA May-Jun Change %',
              'NSA Aug 2026','NSA Jul 2026','NSA Jun 2026','NSA Aug 2025','NSA YTD 2026','NSA YTD 2025','NSA YTD Change %'])
    for row in rows:s.append(row)
    s.append([])
    for i in range(91,107):s.append([m3[i]])
    p=ROOT/'datasets/xlsx/manufacturing/input/m3_shipments_202608.xlsx';p.parent.mkdir(parents=True,exist_ok=True);book.save(p);book.close()
    total=rows[0];durable=next(r for r in rows if r[0]=='Durable goods industries'); nondurable=next(r for r in rows if r[0]=='Nondurable goods industries')
    values=dict(total_shipments=total[1],monthly_growth=total[1]/total[2]-1,durable_share=durable[1]/total[1],nondurable_shipments=nondurable[1])
    assumptions=[dict(name=n,assumptions=dict(shipments_multiplier=v)) for n,v in [('Base',1),('Expansion',1.03),('Contraction',.97)]]
    scenario=dict(default='Base',dimensions=[dict(id='shipments_multiplier',meaning='Hypothetical shipments multiplier, not an official forecast')],rows=assumptions,
                  projections=[dict(id='projected_shipments',meaning='Total adjusted shipments scaled by the selected multiplier',unit='USD million')])
    freeze('manufacturing',p,[metric('total_shipments','Total shipments','All manufacturing industries seasonally adjusted August 2026 shipments. Do not sum overlapping industry hierarchy rows.','USD million'),
        metric('monthly_growth','Monthly change','Relative change in all-manufacturing adjusted shipments July to August 2026 computed from the levels.','fraction'),
        metric('durable_share','Durable goods share','Durable-goods August 2026 adjusted shipments as a fraction of all manufacturing shipments.','fraction'),
        metric('nondurable_shipments','Nondurable shipments','Published nondurable-goods August 2026 seasonally adjusted shipments.','USD million')],values,scenario,
        {r['name']:dict(projected_shipments=values['total_shipments']*r['assumptions']['shipments_multiplier']) for r in assumptions},
        [dict(id='durability_mix',meaning='Compare August 2026 seasonally adjusted durable versus nondurable shipments, excluding subordinate rows',unit='USD million'),
         dict(id='shipment_trend',meaning='All-manufacturing seasonally adjusted shipments June, July, August 2026 in chronological order',unit='USD million')],
        dict(task='Analyze Census M3 August 2026 Table 1 shipments',input_class='official_table_transcription',native_xlsx=False,
             official_url='https://www.census.gov/manufacturing/m3/prel/pdf/table1p.pdf',official_xlsx_url='https://www.census.gov/manufacturing/m3/prel/table1p.xlsx',
             actual_download_source='Official census.gov PDF Table 1 via web-tool text extraction',source_filename='table1p.pdf',release='August 2026 Full Report, CB26-150, 2026-10-02',download_date='2026-10-09',
             source_evidence='sources/provenance-audit/m3-table1-web.txt',source_evidence_sha256=sha256(ROOT/'sources/provenance-audit/m3-table1-web.txt'),
             observations=len(rows),synthetic_expansion=False,conversion='All 71 Table 1 rows (including aggregate, exclusion and subgroup rows), 13 measure columns and notes retained; wrapped labels joined. No native XLSX or PDF byte identity claimed.',
             limitations=['Native Census XLSX downloads return HTTP 403; official-table fallback, not native XLSX.']))


def build_hr():
    source=ROOT/'sources/xlsx/national_M2025_dl.xlsx'
    book=load_workbook(source,read_only=True,data_only=True)
    rows=list(book['national_M2025_dl'].iter_rows(values_only=True)); header=rows[0];data=[dict(zip(header,r)) for r in rows[1:]]
    record=next(r for r in data if r['OCC_CODE']=='00-0000' and r['NAICS']=='000000' and str(r['OWN_CODE'])=='1235')
    major=[r for r in data if r['O_GROUP']=='major' and r['NAICS']=='000000' and str(r['OWN_CODE'])=='1235'];book.close()
    values=dict(total_employment=record['TOT_EMP'],annual_mean_wage=record['A_MEAN'],hourly_median_wage=record['H_MEDIAN'])
    rows=[dict(name=n,assumptions=dict(wage_multiplier=v)) for n,v in [('Base',1),('Increase',1.05)]]
    scenario=dict(default='Base',dimensions=[dict(id='wage_multiplier',meaning='Illustrative mean-wage multiplier; no new employment records')],rows=rows,
                  projections=[dict(id='projected_mean_wage',meaning='Published all-occupations annual mean wage scaled by selected wage multiplier',unit='USD per year')])
    freeze('hr',source,[metric('total_employment','National employment','Published cross-industry, all-ownerships, all-occupations national employment total. Do not sum aggregate and detailed records together.','employees'),
        metric('annual_mean_wage','Annual mean wage','Published annual mean wage for that all-occupations national total; do not calculate an unweighted occupation mean.','USD per year'),
        metric('hourly_median_wage','Hourly median wage','Published hourly median wage for that same national total; do not average occupation medians.','USD per hour')],values,scenario,
        {r['name']:dict(projected_mean_wage=values['annual_mean_wage']*r['assumptions']['wage_multiplier']) for r in rows},
        [dict(id='major_employment',meaning='Employment for the five largest cross-industry all-ownership major occupational groups, ranked descending; exclude totals and detailed/broad/minor rows',unit='employees')],
        dict(task='Analyze the frozen May 2025 OEWS national workbook',input_class='cached_native_spreadsheet',native_xlsx=True,official_byte_identity_verified=False,
             official_url='https://www.bls.gov/oes/special-requests/oesm25nat.zip',
             actual_download_source='https://raw.githubusercontent.com/hack4rva/richmond-ai-impact-analysis/main/data/oesm25nat/national_M2025_dl.xlsx',
             source_filename=source.name,release='May 2025 OEWS estimates',download_date='2026-10-09 (cache first acquired; official retry same day returned 403)',
             synthetic_expansion=False,conversion='None. Unmodified cache bytes; all worksheets, occupation hierarchy, suppression markers and formatting retained.',
             verification_evidence='reports/hr-sba-source-reliability-v1.json',
             verified_fields='825 detailed seed occupations: TOT_EMP and A_MEAN; 767 published H_MEAN/H_MEDIAN values against BLS official news table. Formal KPI cells will be separately corroborated.',
             unverified_fields=['Official ZIP/workbook byte identity','A_MEDIAN and percentile columns','All 1401 original records, including non-seed aggregates and other ownership records'],
             limitations=['A GitHub cache is not an official download. The source limits remain explicit; no silent substitution when the official file becomes accessible.']))


def migrate_tlc():
    root=ROOT/'datasets/xlsx/tlc';old=read_json(ROOT/'legacy/runner/cases/xlsx/manifest.json');values=read_json(ROOT/'legacy/runner/cases/xlsx/expected.json')['cached_values']
    metrics=[metric('total_trips','Total trips','Trip count represented by the original workbook full-population summaries, rather than counting only sample rows.','trips'),
        metric('fare_revenue','Fare revenue','Total fare revenue for the cleaned full population.','USD'),
        metric('average_fare','Average fare','Overall average fare from the original population summaries.','USD per trip'),
        metric('removal_rate','Removal rate','Fraction of original trips removed during cleaning.','fraction'),
        metric('reconciliation_status','Reconciliation status','PASS only when both original population consistency checks indicate consistency; otherwise REVIEW.','text')]
    rows=[dict(name=n,assumptions=dict(fare_multiplier=f,trip_multiplier=t)) for n,f,t in [('Base',1,1),('Upside',1.08,1.05),('Downside',.93,.95)]]
    scenario=dict(default='Base',dimensions=[dict(id='fare_multiplier',meaning='Hypothetical average-fare multiplier'),dict(id='trip_multiplier',meaning='Hypothetical trip-volume multiplier')],rows=rows,
        projections=[dict(id='projected_trips',meaning='Population trip count scaled by selected trip-volume multiplier',unit='trips'),
                     dict(id='projected_fare',meaning='Population fare revenue scaled by both selected fare and trip multipliers',unit='USD')])
    freeze('tlc',root/'input/monthly_operations_template.xlsx',metrics,dict(zip([m['id'] for m in metrics],[values[c] for c in ('B5','B6','B7','B8','B9')])),scenario,
        {r['name']:dict(projected_trips=values['B5']*r['assumptions']['trip_multiplier'],projected_fare=values['B6']*r['assumptions']['fare_multiplier']*r['assumptions']['trip_multiplier']) for r in rows},
        [dict(id='daily_trips',meaning='Daily full-population trip count trend in date order',unit='trips'),dict(id='payment_mix',meaning='Payment-method full-population trip count mix',unit='trips')],
        dict(task=old['task'],input_class='benchmark_generated_record_sample',native_xlsx=False,synthetic_expansion=False,
             official_url='https://www.nyc.gov/site/tlc/about/tlc-trip-record-data.page',original_data_url='https://d37ci6vzurychx.cloudfront.net/trip-data/yellow_tripdata_2024-01.parquet',
             source_filename='yellow_tripdata_2024-01.parquet',release='January 2024 yellow taxi records',download_date='Historical import; precise acquisition date not recorded',
             conversion='Frozen historical 100000-real-record sample workbook with full-population aggregation sheets; generated XLSX, not TLC-published spreadsheet. No synthetic shards.',
             limitations=['Original upstream byte-level download evidence is historical and incomplete; frozen workbook hash and legacy artifacts are available.']))


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--initialize-v2',action='store_true',required=True);parser.parse_args()
    for path in (ROOT/'datasets/xlsx').glob('*/manifest.json'):
        if read_json(path).get('schema_version')=='office-dataset-v2':
            raise ValueError('suite v2 is already frozen; migration cannot overwrite any existing input')
    build_tables();build_hr();migrate_tlc()
    print('Frozen 4 XLSX datasets without synthetic expansion. Official-table fallbacks are explicitly labelled.')


if __name__=='__main__':main()
